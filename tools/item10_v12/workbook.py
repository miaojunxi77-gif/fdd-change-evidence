import json,zipfile
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill,Font,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from .config import FREEZE,VERSION
from .outputs import BRAND_FIELDS
DARK='173F35';HEADER='2F6757';SUB='DDEBE5';WHITE='FFFFFF';thin=Side(style='thin',color='D6DED9')
def title(ws,text,n):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n);c=ws.cell(1,1,text);c.fill=PatternFill('solid',fgColor=DARK);c.font=Font(color=WHITE,bold=True,size=14);c.alignment=Alignment(vertical='center');ws.row_dimensions[1].height=24
def headers(ws,row,vals):
    for j,v in enumerate(vals,1):
        c=ws.cell(row,j,v);c.fill=PatternFill('solid',fgColor=HEADER);c.font=Font(color=WHITE,bold=True);c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(bottom=thin)
    ws.auto_filter.ref=f"A{row}:{get_column_letter(len(vals))}{row}"
def widths(ws,d):
    for c,w in d.items():ws.column_dimensions[c].width=w
def add(ws,start,data,fields=None):
    for i,r in enumerate(data,start):
        vals=[r.get(f,'') for f in fields] if fields else r
        for j,v in enumerate(vals,1):ws.cell(i,j,v)
def build(data,outdir):
    outdir=Path(outdir);outdir.mkdir(parents=True,exist_ok=True);c=data['counts'];wb=Workbook();wb.remove(wb.active)
    ws=wb.create_sheet('Read Me');title(ws,'Item 10 Financing — Advisor Package (Production V1.2)',8);rm=[('Production freeze',FREEZE),('Primary unit','Canonical franchise offering/system × filing year after chronology and offering-scope QC'),('Reconciled brand-years',c['all']),('Paper-ready brand-years',c['paper']),('Unresolved',c['unresolved']),('Paper-ready coverage',c['paper']/c['all']),('Latest snapshot','One row per V1.2 identity, using the latest paper-ready year'),('Consecutive within-system pairs',c['pairs']),('RiskScore changes',c['changes']),('Chronology rule','Use a later amendment/version only when supersession is unambiguous within the same offering scope.'),('Offering rule','Split distinct programs/offering types instead of forcing them into one system-year score.'),('State rule','Verified state-specific Item 10 differences remain NA in the national system-year panel.'),('Original evidence','Original FDD/PDF remains authoritative; extracted text and quotes are audit aids.')];headers(ws,3,['Field','Description']);add(ws,4,rm);ws['B9'].number_format='0.00%';widths(ws,{'A':30,'B':100});ws.freeze_panes='A4'
    ws=wb.create_sheet('Main Results');title(ws,'Main Results — Production V1.2',7);headers(ws,3,['Metric','V1.1','V1.2','Change','','Outcome','V1.2 share']);mm=[('Reconciled brand-years',13222,c['all']),('Paper-ready brand-years',13067,c['paper']),('Unresolved score conflicts',155,c['unresolved']),('Consecutive within-brand pairs',7013,c['pairs']),('RiskScore changes',341,c['changes']),('0 → support',130,c['zero_to_support']),('Support → 0',144,c['support_to_zero'])]
    for i,(m,a,b) in enumerate(mm,4):ws.cell(i,1,m);ws.cell(i,2,a);ws.cell(i,3,b);ws.cell(i,4,b-a)
    oc=[('Any support (score ≥ 1)',c['any']/c['paper']),('Risk-bearing support (score ≥ 2)',c['risk']/c['paper']),('Material credit risk (score ≥ 3)',c['material']/c['paper']),('Direct funded credit (score = 4)',c['direct']/c['paper']),('RiskScore change rate',c['changes']/c['pairs'])]
    for i,(m,v) in enumerate(oc,4):ws.cell(i,6,m);ws.cell(i,7,v);ws.cell(i,7).number_format='0.00%'
    ws.merge_cells('A13:G13');ws['A13']='Interpretation';ws['A13'].fill=PatternFill('solid',fgColor=SUB);ws['A13'].font=Font(bold=True);ws.merge_cells('A14:G16');ws['A14']=f"Production V1.2 contains {c['all']:,} reconciled offering/system-year observations, of which {c['paper']:,} are paper-ready ({c['paper']/c['all']:.2%}). Financing support appears in {c['any']/c['paper']:.2%}; {c['risk']/c['paper']:.2%} involve risk-bearing support, {c['material']/c['paper']:.2%} material credit risk, and {c['direct']/c['paper']:.2%} direct funded credit. Among {c['pairs']:,} consecutive within-identity pairs, {c['changes']:,} change RiskScore ({c['changes']/c['pairs']:.2%}).";ws['A14'].alignment=Alignment(wrap_text=True,vertical='top');widths(ws,{'A':35,'B':13,'C':13,'D':13,'E':3,'F':35,'G':16});ws.freeze_panes='A4'
    ws=wb.create_sheet('Latest Franchisor');title(ws,'Latest paper-ready observation per V1.2 franchise offering/system',12);lf=['canonical_chain_id','canonical_chain_name','latest_paper_ready_year','final_risk_score','any_support','risk_bearing_ge2','material_risk_ge3','direct_funded_eq4','final_forms','providers','purposes','evidence_quote'];headers(ws,3,['Canonical Chain ID','Franchise Offering/System','Latest Year','Risk Score','Any Support','Risk-bearing ≥2','Material Risk ≥3','Direct Funded =4','Financing Forms','Providers','Purposes','Evidence Quote']);add(ws,4,data['latest'],lf);widths(ws,{'A':20,'B':40,'C':12,'D':11,'E':12,'F':15,'G':15,'H':15,'I':42,'J':26,'K':30,'L':75});ws.freeze_panes='C4'
    for x in ws['L'][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    ws=wb.create_sheet('Brand-Year Evidence');title(ws,'Production V1.2 brand-year evidence panel',16);headers(ws,3,['Chain ID','Franchise Offering/System','Year','Risk Score','Status','Paper Ready','Evidence Versions','Unique Item 10 Versions','Sources','Financing Forms','Providers','Purposes','Evidence Quote','Full Text','V1.2 QC Status','V1.2 Resolution Rule']);add(ws,4,data['brand'],BRAND_FIELDS);widths(ws,{'A':20,'B':42,'C':10,'D':11,'E':30,'F':12,'G':14,'H':18,'I':28,'J':44,'K':27,'L':30,'M':75,'N':12,'O':40,'P':75});ws.freeze_panes='C4'
    for col in ('M','P'):
        for x in ws[col][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    ws.conditional_formatting.add(f'D4:D{ws.max_row}',ColorScaleRule(start_type='min',start_color='E8F1ED',mid_type='percentile',mid_value=50,mid_color='F6E8B1',end_type='max',end_color='F3C1B8'))
    ws=wb.create_sheet('Within-Firm Changes');title(ws,'Production V1.2 consecutive-year RiskScore changes',14);wf=['canonical_chain_id','canonical_chain_name','from_year','to_year','from_score','to_score','score_change','change_type','from_forms','to_forms','from_providers','to_providers','from_evidence_quote','to_evidence_quote'];headers(ws,3,['Chain ID','Franchise Offering/System','From Year','To Year','From Score','To Score','Δ Score','Change Type','Before Forms','After Forms','Before Providers','After Providers','Before Evidence','After Evidence']);add(ws,4,data['changes'],wf);widths(ws,{'A':20,'B':42,'C':11,'D':11,'E':11,'F':11,'G':10,'H':20,'I':38,'J':38,'K':26,'L':26,'M':70,'N':70});ws.freeze_panes='C4'
    for col in ('M','N'):
        for x in ws[col][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    for nm,txt,key in [('QC Audit','All 155 V1.1 unresolved conflicts — V1.2 disposition','audit'),('Scope Splits','25 offering / identity scope cases','scope'),('Chronology Resolved','19 conflicts resolved by unambiguous chronology','chronology')]:
        ws=wb.create_sheet(nm);title(ws,txt,10);af=['original_index','brand','year','old_id','action','new_id','old_score','new_score','outcome','reason'];headers(ws,3,['V1.1 #','Brand','Year','Old ID','V1.2 Action','V1.2 ID / split targets','Old Score','New Score','Outcome','Reason']);add(ws,4,data[key],af);widths(ws,{'A':9,'B':37,'C':9,'D':28,'E':42,'F':55,'G':11,'H':18,'I':30,'J':85});ws.freeze_panes='D4'
        for col in ('E','F','I','J'):
            for x in ws[col][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    ws=wb.create_sheet('Unresolved');title(ws,'113 unresolved observations retained as NA in Production V1.2',8);headers(ws,3,['Chain ID','Franchise Offering/System','Year','Sources','Conflicting Forms','V1.2 QC Status','Reason','Evidence Excerpt']);ur=[{'id':r['chainId'],'brand':r['brand'],'year':r['year'],'sources':' | '.join(r.get('sources') or []),'forms':' | '.join(r.get('forms') or []),'qc':r.get('v1_2_qc_status',''),'reason':r.get('v1_2_resolution_rule',''),'evidence':r.get('evidenceQuote','')} for r in data['unresolved']];add(ws,4,ur,['id','brand','year','sources','forms','qc','reason','evidence']);widths(ws,{'A':22,'B':42,'C':9,'D':30,'E':45,'F':45,'G':80,'H':70});ws.freeze_panes='C4'
    for col in ('D','E','F','G','H'):
        for x in ws[col][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    ws=wb.create_sheet('QC Summary');title(ws,'Original 155-conflict chronology / identity QC classification',3);headers(ws,3,['QC Class','N original conflicts','V1.2 policy']);add(ws,4,data['qc_summary'],['qc_class','n_original_conflicts','v1_2_policy']);widths(ws,{'A':48,'B':20,'C':95});ws.freeze_panes='A4'
    for x in ws['C'][3:]:x.alignment=Alignment(wrap_text=True,vertical='top')
    for ws in wb.worksheets:ws.sheet_view.showGridLines=False
    xlsx=outdir/'Item10_Financing_Advisor_Package_V1_2.xlsx';wb.save(xlsx)
    check=load_workbook(xlsx,read_only=True,data_only=False);assert check.sheetnames==['Read Me','Main Results','Latest Franchisor','Brand-Year Evidence','Within-Firm Changes','QC Audit','Scope Splits','Chronology Resolved','Unresolved','QC Summary'];assert check['Brand-Year Evidence'].max_row==c['all']+3;assert check['Within-Firm Changes'].max_row==c['changes']+3;assert check['Unresolved'].max_row==c['unresolved']+3;check.close()
    readme=f"""Item 10 Financing — Production V1.2\nFreeze: {FREEZE}\n\nCore results\n- Reconciled offering/system-years: {c['all']:,}\n- Paper-ready: {c['paper']:,}\n- Unresolved: {c['unresolved']:,}\n- Coverage: {c['paper']/c['all']:.2%}\n- Any support: {c['any']/c['paper']:.2%}\n- Risk-bearing >=2: {c['risk']/c['paper']:.2%}\n- Material risk >=3: {c['material']/c['paper']:.2%}\n- Direct funded =4: {c['direct']/c['paper']:.2%}\n- Consecutive pairs: {c['pairs']:,}\n- RiskScore changes: {c['changes']:,} ({c['changes']/c['pairs']:.2%})\n\nQC policy\n1. Apply chronology only when a later amendment/version unambiguously supersedes an earlier version within the same offering.\n2. Split distinct offerings/programs before scoring.\n3. Preserve verified state-specific substantive differences as NA in the national offering/system-year panel.\n4. Do not force cases whose identity, year, or chronology remains ambiguous.\n\nThe original FDD/PDF remains the authoritative source.\n""";(outdir/'README_V1_2.txt').write_text(readme,encoding='utf-8');manifest={'version':VERSION,'freeze':FREEZE,'counts':c,'score_counts':dict(data['score_counts'])};(outdir/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    zpath=outdir.parent/'Item10_Financing_Production_V1_2_Release.zip'
    with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(outdir.iterdir()):
            if p.is_file():z.write(p,arcname=p.name)
    with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
    return xlsx,zpath
